"""Gera a apresentação institucional do Área Comum (PPTX) com as telas reais."""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(r"D:\Users\vivis.DESKTOP-RN5T169\Downloads\Deepseek Harness\area-comum")
ASSETS = ROOT / "apresentacao" / "assets"
PPTX = ROOT / "apresentacao" / "Apresentacao-Area-Comum-Sao-Cristovao.pptx"

INK = RGBColor(0x12, 0x2F, 0x27)
INK_2 = RGBColor(0x1B, 0x40, 0x36)
CREAM = RGBColor(0xF7, 0xF2, 0xE8)
CARD = RGBColor(0xFD, 0xFB, 0xF7)
LINE = RGBColor(0xE7, 0xDD, 0xCB)
GOLD = RGBColor(0xD9, 0xA4, 0x41)
GOLD_LIGHT = RGBColor(0xEA, 0xC4, 0x69)
GOLD_PALE = RGBColor(0xF9, 0xEF, 0xCE)
GOLD_DARK = RGBColor(0x8A, 0x5F, 0x1E)
MUTED = RGBColor(0x64, 0x74, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xD6, 0x1F, 0x5C)
BLUE = RGBColor(0x14, 0x6F, 0xAE)

SERIF = "Georgia"
SANS = "Segoe UI"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def rgb(hex_color):
    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def new_slide(background=CREAM):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = background
    return slide


def add_rect(slide, x, y, w, h, fill, border=None, radius=None, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try:
            shape.adjustments[0] = min(0.5, radius)
        except Exception:
            pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = border
        shape.line.width = Pt(0.9)
    if transparency:
        # aplica alpha no preenchimento
        from lxml import etree
        from pptx.oxml.ns import qn

        spPr = shape._element.spPr
        solid = spPr.find(qn("a:solidFill"))
        if solid is not None:
            clr = solid.find(qn("a:srgbClr"))
            if clr is not None:
                alpha = etree.SubElement(clr, qn("a:alpha"))
                alpha.set("val", str(int(transparency * 1000)))
    return shape


def add_line(slide, x, y, w, color=LINE, height=0.012):
    add_rect(slide, x, y, w, height, color)


def add_circle(slide, x, y, d, fill, border=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = border
        shape.line.width = Pt(0.8)
    return shape


def add_text(slide, x, y, w, h, paragraphs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=5, line_spacing=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for index, runs in enumerate(paragraphs):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(space_after)
        paragraph.line_spacing = line_spacing
        for text, opts in runs:
            run = paragraph.add_run()
            run.text = text
            font = run.font
            font.name = opts.get("font", SANS)
            font.size = Pt(opts.get("size", 12))
            font.bold = opts.get("bold", False)
            font.italic = opts.get("italic", False)
            font.color.rgb = opts.get("color", INK)
    return box


def add_pill(slide, x, y, w, h, text, fill, text_color, border=None, size=9, bold=True):
    add_rect(slide, x, y, w, h, fill, border=border, radius=0.5)
    add_text(slide, x, y, w, h, [[(text, {"size": size, "bold": bold, "color": text_color})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def add_badge(slide, x, y, d=0.52):
    slide.shapes.add_picture(str(ROOT / "public" / "brand" / "emblem-sc-badge.png"),
                             Inches(x), Inches(y), Inches(d), Inches(d))


def add_image_fit(slide, path, x, y, w, h, align="center", valign="middle"):
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    dx = x + (w - dw) / 2 if align == "center" else x
    dy = y + (h - dh) / 2 if valign == "middle" else y
    return slide.shapes.add_picture(str(path), Inches(dx), Inches(dy), Inches(dw), Inches(dh))


def footer(slide, page, dark=False):
    color = RGBColor(0x7E, 0x8E, 0x86) if dark else RGBColor(0x9A, 0xA6, 0x9F)
    line_color = RGBColor(0x3A, 0x55, 0x4B) if dark else LINE
    add_line(slide, 0.65, 7.08, 12.03, line_color)
    add_text(slide, 0.65, 7.15, 7.5, 0.2,
             [[("Condomínio São Cristóvão · Área Comum", {"size": 7.5, "color": color})]])
    add_text(slide, 11.65, 7.15, 1.03, 0.2,
             [[(f"{page:02d}", {"size": 7.5, "color": color})]], align=PP_ALIGN.RIGHT)


def header(slide, eyebrow, title, subtitle=None, title_size=31, width=11.4, subtitle_width=10.3):
    add_text(slide, 0.65, 0.5, 8.5, 0.25,
             [[(eyebrow, {"size": 8.5, "bold": True, "color": GOLD_DARK})]])
    add_text(slide, 0.65, 0.83, width, 1.25, title, space_after=0, line_spacing=0.96)
    if subtitle:
        add_text(slide, 0.65, 2.05, subtitle_width, 0.55,
                 [[(subtitle, {"size": 10.5, "color": MUTED})]])
    add_badge(slide, 12.55, 0.42, 0.46)


def number_badge(slide, x, y, n, d=0.34):
    add_circle(slide, x, y, d, GOLD_PALE)
    add_text(slide, x, y, d, d, [[(f"{n:02d}", {"size": 8.5, "bold": True, "color": GOLD_DARK})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------------------
# 01 — Capa
# ---------------------------------------------------------------------------
slide = new_slide(INK)
slide.shapes.add_picture(str(ASSETS / "watermark.png"), Inches(7.3), Inches(0.6), Inches(6.7), Inches(6.7))
add_badge(slide, 0.68, 0.55, 0.58)
add_text(slide, 1.42, 0.62, 5.4, 0.7,
         [[("São Cristóvão", {"size": 15, "bold": True, "color": CREAM})],
          [("CONDOMÍNIO · ÁREA COMUM", {"size": 7.5, "color": RGBColor(0xB9, 0xCF, 0xC4)})]],
         space_after=1)
add_text(slide, 0.7, 2.28, 7.0, 0.3,
         [[("PROPOSTA DE IMPLANTAÇÃO · OUTUBRO DE 2026", {"size": 8.5, "bold": True, "color": GOLD_LIGHT})]])
add_text(slide, 0.68, 2.65, 7.2, 1.15,
         [[("Área Comum", {"font": SERIF, "size": 54, "color": CREAM})]])
add_text(slide, 0.7, 3.82, 7.0, 1.35,
         [[("Reservas do condomínio,", {"font": SERIF, "size": 24, "color": CREAM})],
          [("em sintonia.", {"font": SERIF, "size": 30, "italic": True, "color": GOLD_LIGHT})]],
         space_after=2)
add_text(slide, 0.7, 5.42, 5.8, 0.9,
         [[("Agenda, regras, aprovações, pagamentos e avisos em uma única plataforma — "
            "com cada pessoa vendo e fazendo exatamente o que o seu papel permite.", {"size": 10.5, "color": RGBColor(0xD7, 0xE4, 0xDD)})]],
         line_spacing=1.25)
for index, (label, width) in enumerate([("Morador", 1.02), ("Síndico", 1.12), ("Portaria", 1.15)]):
    add_pill(slide, 0.7 + sum([1.02, 1.12][:index]) + 0.12 * index, 6.32, width, 0.32,
             label, INK_2, CREAM, border=RGBColor(0x3A, 0x55, 0x4B), size=8.5)
add_line(slide, 0.7, 6.92, 11.95, RGBColor(0x3A, 0x55, 0x4B))
add_text(slide, 0.7, 7.02, 6.0, 0.2,
         [[("Condomínio São Cristóvão", {"size": 7.5, "color": RGBColor(0x9F, 0xB3, 0xA8)})]])
add_text(slide, 8.0, 7.02, 4.65, 0.2,
         [[("Sistema web · Next.js + SQLite", {"size": 7.5, "color": RGBColor(0x9F, 0xB3, 0xA8)})]],
         align=PP_ALIGN.RIGHT)
notes(slide, "Apresentação institucional do Área Comum para o Condomínio São Cristóvão. "
             "O objetivo é mostrar a solução, a experiência dos três perfis e as regras que ficam "
             "garantidas no servidor.")

# ---------------------------------------------------------------------------
# 02 — Problema
# ---------------------------------------------------------------------------
slide = new_slide()
header(slide, "DIAGNÓSTICO", [
    [("O caderno da portaria", {"font": SERIF, "size": 31, "color": INK})],
    [("não dá mais conta.", {"font": SERIF, "size": 31, "italic": True, "color": GOLD_DARK})],
], "Quatro áreas, dezenas de moradores e regras que hoje vivem na memória de cada um.")

problems = [
    ("01", "Briga por horário", "Duas famílias acreditam ter reservado o mesmo espaço no mesmo horário."),
    ("02", "Agenda invisível", "O morador não sabe se a quadra está livre, quem reservou ou quando será liberada."),
    ("03", "Regras sem registro", "Aprovações, bloqueios e multas ficam informais — e difíceis de cobrar depois."),
]
positions = [0.65, 4.53, 8.41]
for (number, title, text), x in zip(problems, positions):
    add_rect(slide, x, 2.75, 3.75, 2.35, CARD, border=LINE, radius=0.05)
    number_badge(slide, x + 0.28, 2.98, int(number))
    add_text(slide, x + 0.28, 3.48, 3.1, 0.35,
             [[(title, {"size": 13, "bold": True, "color": INK})]])
    add_text(slide, x + 0.28, 3.88, 3.1, 1.0,
             [[(text, {"size": 9.5, "color": MUTED})]], line_spacing=1.2)

add_rect(slide, 0.65, 5.55, 12.03, 1.05, INK, radius=0.05)
add_text(slide, 1.0, 5.72, 10.4, 0.75,
         [[("“Hoje isso é feito num caderno na portaria e sempre dá briga por horário.”",
            {"font": SERIF, "size": 16, "italic": True, "color": CREAM})],
          [("— a realidade que o Área Comum resolve", {"size": 8.5, "color": GOLD_LIGHT})]],
         space_after=3)
footer(slide, 2)
notes(slide, "O problema não é apenas estético: a falta de uma fonte única de verdade gera conflito, "
             "retrabalho na portaria e regras que ninguém consegue auditar.")

# ---------------------------------------------------------------------------
# 03 — Solução
# ---------------------------------------------------------------------------
slide = new_slide()
header(slide, "A SOLUÇÃO", [
    [("Um sistema de verdade —", {"font": SERIF, "size": 31, "color": INK})],
    [("não só telas bonitas.", {"font": SERIF, "size": 31, "italic": True, "color": GOLD_DARK})],
], "Agenda, regras, aprovações, pagamentos e avisos no mesmo lugar, com permissões por perfil.",
    subtitle_width=5.35)

items = [
    ("Agenda e calendário das 4 áreas", "Visão mensal, por dia e por espaço, com bloqueios de manutenção."),
    ("Regras e aprovações automáticas", "Conflito de horário, antecedência, limite de convidados e 2 reservas futuras."),
    ("Pagamentos e notificações", "Multas com PIX/cartão simulado, comprovante e central de avisos."),
    ("Permissões garantidas no servidor", "Cada papel vê e faz somente o que lhe cabe — em cada API e consulta."),
]
for index, (title, text) in enumerate(items):
    y = 2.72 + index * 0.92
    add_circle(slide, 0.68, y + 0.02, 0.24, GOLD_PALE)
    add_text(slide, 0.68, y + 0.02, 0.24, 0.24,
             [[("✓", {"size": 9, "bold": True, "color": GOLD_DARK})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 1.05, y - 0.02, 4.8, 0.28,
             [[(title, {"size": 12, "bold": True, "color": INK})]])
    add_text(slide, 1.05, y + 0.27, 4.9, 0.5,
             [[(text, {"size": 9, "color": MUTED})]], line_spacing=1.15)

add_image_fit(slide, ASSETS / "visao-geral.png", 6.35, 1.35, 6.45, 5.45)
add_pill(slide, 6.4, 6.78, 1.15, 0.26, "Visão do morador", GOLD_PALE, GOLD_DARK, size=7.5)
add_pill(slide, 7.62, 6.78, 1.7, 0.26, "Dados reais de demonstração", CARD, MUTED, border=LINE, size=7.5)
footer(slide, 3)
notes(slide, "A proposta entrega uma plataforma completa: a interface bonita é só a ponta. "
             "As regras e permissões são validadas no servidor, independente do que aparece na tela.")

# ---------------------------------------------------------------------------
# 04 — Jornada do morador
# ---------------------------------------------------------------------------
slide = new_slide()
header(slide, "EXPERIÊNCIA DO MORADOR", [
    [("Do celular para a agenda,", {"font": SERIF, "size": 31, "color": INK})],
    [("em poucos toques.", {"font": SERIF, "size": 31, "italic": True, "color": GOLD_DARK})],
], "O morador consulta a disponibilidade, reserva, paga multas e recebe avisos sem passar na portaria.")

cards = [
    ("1", "Visão geral", "Próxima reserva, calendário do mês e situação do apartamento.", "visao-geral.png"),
    ("2", "Agenda", "Escolhe data, área e horário; as regras são conferidas na hora.", "agenda.png"),
    ("3", "Minhas reservas", "Acompanha solicitações, cancela e recebe notificações.", "minhas.png"),
]
for index, (number, title, caption, image) in enumerate(cards):
    x = 0.65 + index * 4.18
    add_rect(slide, x, 2.35, 3.93, 4.15, CARD, border=LINE, radius=0.04)
    add_image_fit(slide, ASSETS / image, x + 0.12, 2.5, 3.69, 2.55)
    number_badge(slide, x + 0.2, 5.12, int(number), 0.3)
    add_text(slide, x + 0.6, 5.09, 3.0, 0.3,
             [[(title, {"size": 12, "bold": True, "color": INK})]])
    add_text(slide, x + 0.2, 5.48, 3.5, 0.85,
             [[(caption, {"size": 8.8, "color": MUTED})]], line_spacing=1.2)

add_rect(slide, 0.65, 6.62, 12.03, 0.3, INK, radius=0.5)
add_text(slide, 0.88, 6.66, 11.6, 0.22,
         [[("Salão e churrasqueira passam pela aprovação do síndico · quadra e sala confirmam na hora.",
            {"size": 8, "color": CREAM})]])
footer(slide, 4)
notes(slide, "A jornada do morador foi desenhada para reduzir idas à portaria: ele vê a agenda, "
             "faz a reserva, acompanha a aprovação e resolve multas pelo próprio celular.")

# ---------------------------------------------------------------------------
# 05 — Calendário
# ---------------------------------------------------------------------------
slide = new_slide()
header(slide, "AGENDA", [
    [("Calendário mensal e", {"font": SERIF, "size": 31, "color": INK})],
    [("agenda por área.", {"font": SERIF, "size": 31, "italic": True, "color": GOLD_DARK})],
])
add_text(slide, 8.18, 1.38, 4.5, 0.55,
         [[("Uma leitura rápida do mês inteiro e o detalhe de cada espaço.", {"size": 9.5, "color": MUTED})]],
         line_spacing=1.2)
add_image_fit(slide, ASSETS / "agenda.png", 0.5, 1.95, 7.15, 5.0)

bullets = [
    ("Calendário do mês", "Reservas confirmadas e pendentes aparecem com pontos coloridos por área."),
    ("Bloqueios de manutenção", "O síndico marca o dia; a agenda e o formulário refletem imediatamente."),
    ("Agenda de 21 dias por área", "O morador vê os próximos dias, horários livres e regras específicas."),
]
for index, (title, text) in enumerate(bullets):
    y = 2.15 + index * 1.23
    add_circle(slide, 8.18, y + 0.02, 0.24, GOLD_PALE)
    add_text(slide, 8.18, y + 0.02, 0.24, 0.24,
             [[("✓", {"size": 9, "bold": True, "color": GOLD_DARK})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 8.55, y - 0.03, 4.2, 0.28,
             [[(title, {"size": 12, "bold": True, "color": INK})]])
    add_text(slide, 8.55, y + 0.27, 4.2, 0.78,
             [[(text, {"size": 9, "color": MUTED})]], line_spacing=1.2)

add_rect(slide, 8.2, 5.72, 4.45, 0.95, INK, radius=0.05)
add_text(slide, 8.45, 5.86, 3.95, 0.68,
         [[("Uma fonte de verdade", {"size": 11, "bold": True, "color": CREAM})],
          [("para moradores, síndico e portaria.", {"size": 8.5, "color": RGBColor(0xB9, 0xCF, 0xC4)})]],
         space_after=2)
footer(slide, 5)
notes(slide, "O calendário mensal é a visão de conjunto. A página de cada área aprofunda os próximos "
             "21 dias, mostra as regras e permite reservar sem sair do contexto.")

# ---------------------------------------------------------------------------
# 06 — Regras no servidor
# ---------------------------------------------------------------------------
slide = new_slide(INK)
slide.shapes.add_picture(str(ASSETS / "watermark.png"), Inches(9.9), Inches(4.4), Inches(3.6), Inches(3.6))
add_text(slide, 0.65, 0.5, 8.5, 0.25,
         [[("REGRAS DE NEGÓCIO", {"size": 8.5, "bold": True, "color": GOLD_LIGHT})]])
add_text(slide, 0.65, 0.83, 9.8, 1.3,
         [[("As regras valem no servidor.", {"font": SERIF, "size": 32, "color": CREAM})]])
add_text(slide, 0.65, 1.95, 10.2, 0.5,
         [[("A tela orienta. A API e o banco garantem que a regra seja cumprida.", {"size": 10.5, "color": RGBColor(0xB9, 0xCF, 0xC4)})]])

rules = [
    ("01", "Sem sobreposição", "Duas reservas não ocupam o mesmo horário na mesma área."),
    ("02", "Até 2 futuras", "Cada apartamento mantém no máximo duas reservas futuras."),
    ("03", "3 dias + aprovação", "Salão e churrasqueira exigem antecedência e análise do síndico."),
    ("04", "Limite de convidados", "Cada área tem seu próprio teto de pessoas por reserva."),
    ("05", "Multa de R$ 50", "Cancelamento com menos de 24 horas gera multa registrada."),
    ("06", "Multa bloqueia", "Apartamento com multa em aberto não faz nova reserva."),
]
for index, (number, title, text) in enumerate(rules):
    col = index % 3
    row = index // 3
    x = 0.65 + col * 3.9
    y = 2.65 + row * 1.65
    add_rect(slide, x, y, 3.68, 1.45, INK_2, border=RGBColor(0x3A, 0x55, 0x4B), radius=0.05)
    add_text(slide, x + 0.22, y + 0.16, 0.6, 0.3,
             [[(number, {"font": SERIF, "size": 15, "color": GOLD_LIGHT})]])
    add_text(slide, x + 0.78, y + 0.17, 2.7, 0.27,
             [[(title, {"size": 10.5, "bold": True, "color": CREAM})]])
    add_text(slide, x + 0.22, y + 0.62, 3.2, 0.7,
             [[(text, {"size": 8.3, "color": RGBColor(0xC4, 0xD6, 0xCC)})]], line_spacing=1.15)

add_text(slide, 0.65, 6.32, 11.9, 0.5,
         [[("“Cada pessoa só pode ver e fazer o que o papel dela permite — e isso é garantido no servidor.”",
            {"font": SERIF, "size": 13, "italic": True, "color": GOLD_LIGHT})]])
footer(slide, 6, dark=True)
notes(slide, "Este é o diferencial técnico: as regras não dependem de esconder botões. "
             "Elas são verificadas no serviço, na API e no banco, inclusive em chamadas diretas.")

# ---------------------------------------------------------------------------
# 07 — Síndico
# ---------------------------------------------------------------------------
slide = new_slide()
header(slide, "SÍNDICO", [
    [("O síndico no controle,", {"font": SERIF, "size": 31, "color": INK})],
    [("sem planilhas.", {"font": SERIF, "size": 31, "italic": True, "color": GOLD_DARK})],
])
add_text(slide, 7.75, 1.36, 5.15, 0.4,
         [[("Aprovações, áreas, manutenção, multas e uso do mês em um painel só.", {"size": 9.3, "color": MUTED})]],
         line_spacing=1.2)
add_image_fit(slide, ASSETS / "painel.png", 0.45, 1.95, 7.0, 5.0)

add_pill(slide, 7.75, 1.88, 1.15, 0.28, "2 pendentes", GOLD_PALE, GOLD_DARK, size=7.5)
add_pill(slide, 9.0, 1.88, 1.45, 0.28, "R$ 50 em aberto", RGBColor(0xFB, 0xE4, 0xE9), RGBColor(0x9E, 0x1F, 0x45), size=7.5)
add_pill(slide, 10.55, 1.88, 1.95, 0.28, "8 reservas no mês", CARD, MUTED, border=LINE, size=7.5)

items = [
    ("Aprovações", "Aprova ou recusa solicitações com um clique."),
    ("Áreas e regras", "Cadastra espaços, limite de convidados, antecedência e aprovação."),
    ("Manutenção", "Bloqueia datas e o sistema impede reservas naquele dia."),
    ("Multas e uso", "Quita multas e acompanha horas, convidados e ocupação por área."),
]
for index, (title, text) in enumerate(items):
    y = 2.62 + index * 1.0
    add_circle(slide, 7.75, y + 0.02, 0.24, GOLD_PALE)
    add_text(slide, 7.75, y + 0.02, 0.24, 0.24,
             [[("✓", {"size": 9, "bold": True, "color": GOLD_DARK})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 8.12, y - 0.02, 4.4, 0.26,
             [[(title, {"size": 11.5, "bold": True, "color": INK})]])
    add_text(slide, 8.12, y + 0.26, 4.3, 0.55,
             [[(text, {"size": 8.7, "color": MUTED})]], line_spacing=1.15)
footer(slide, 7)
notes(slide, "O síndico ganha uma visão consolidada: aprovações pendentes, regras de cada área, "
             "bloqueios de manutenção, multas e o uso mensal, sem depender de controles paralelos.")

# ---------------------------------------------------------------------------
# 08 — Portaria
# ---------------------------------------------------------------------------
slide = new_slide()
header(slide, "PORTARIA", [
    [("A portaria vê só", {"font": SERIF, "size": 31, "color": INK})],
    [("o que precisa.", {"font": SERIF, "size": 31, "italic": True, "color": GOLD_DARK})],
])
add_text(slide, 7.88, 1.36, 4.7, 0.4,
         [[("Reservas do dia, nome, apartamento e convidados — com entrega e devolução de chave.", {"size": 9.3, "color": MUTED})]],
         line_spacing=1.2)
add_image_fit(slide, ASSETS / "portaria.png", 0.45, 1.95, 7.1, 5.0)

items = [
    ("Somente o dia", "A portaria não acessa a agenda futura nem outros dias."),
    ("Quem e quantos", "Nome do morador, apartamento e número de convidados."),
    ("Chave controlada", "Entrega e devolução registradas com data e hora."),
    ("Privacidade por padrão", "O filtro é aplicado no servidor, não apenas na interface."),
]
for index, (title, text) in enumerate(items):
    y = 2.1 + index * 0.95
    add_circle(slide, 7.88, y + 0.02, 0.24, GOLD_PALE)
    add_text(slide, 7.88, y + 0.02, 0.24, 0.24,
             [[("✓", {"size": 9, "bold": True, "color": GOLD_DARK})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 8.25, y - 0.02, 4.4, 0.26,
             [[(title, {"size": 11.5, "bold": True, "color": INK})]])
    add_text(slide, 8.25, y + 0.25, 4.3, 0.55,
             [[(text, {"size": 8.7, "color": MUTED})]], line_spacing=1.15)

add_rect(slide, 7.88, 6.05, 4.55, 0.62, INK, radius=0.05)
add_text(slide, 8.08, 6.16, 4.15, 0.4,
         [[("“Sem expor a agenda futura nem dados de outros dias.”", {"font": SERIF, "size": 10.5, "italic": True, "color": CREAM})]])
footer(slide, 8)
notes(slide, "A portaria tem uma visão operacional e restrita. O sistema mostra apenas as reservas "
             "do dia e permite registrar a entrega e a devolução da chave.")

# ---------------------------------------------------------------------------
# 09 — Pagamentos e notificações
# ---------------------------------------------------------------------------
slide = new_slide()
header(slide, "PAGAMENTOS & AVISOS", [
    [("Multas e avisos", {"font": SERIF, "size": 31, "color": INK})],
    [("sem ruído.", {"font": SERIF, "size": 31, "italic": True, "color": GOLD_DARK})],
])
add_text(slide, 9.3, 1.4, 3.4, 0.55,
         [[("Pagamento com comprovante e central de avisos para cada evento.", {"size": 9.3, "color": MUTED})]],
         line_spacing=1.2)
add_image_fit(slide, ASSETS / "pagamento.png", 0.5, 1.95, 3.15, 3.2)
add_image_fit(slide, ASSETS / "notificacoes.png", 3.95, 1.95, 5.1, 3.55)

items = [
    ("PIX ou cartão", "Fluxo de demonstração com comprovante e quitação imediata."),
    ("Síndico também quita", "Registro manual no painel para casos presenciais."),
    ("Notificações reais", "Reservas, aprovações, multas, chaves e manutenção."),
    ("Marcar como lidas", "Cada usuário vê somente os próprios avisos."),
]
for index, (title, text) in enumerate(items):
    y = 2.15 + index * 0.85
    add_text(slide, 9.35, y, 3.3, 0.25,
             [[(title, {"size": 10.5, "bold": True, "color": INK})]])
    add_text(slide, 9.35, y + 0.25, 3.35, 0.55,
             [[(text, {"size": 8.3, "color": MUTED})]], line_spacing=1.15)

add_rect(slide, 0.65, 5.72, 12.03, 0.62, GOLD_PALE, radius=0.05)
add_text(slide, 0.9, 5.83, 11.5, 0.4,
         [[("Pagamento simulado nesta demonstração. Em produção, integra-se a um gateway PIX/PSP.",
            {"size": 8.8, "bold": True, "color": GOLD_DARK})]])
footer(slide, 9)
notes(slide, "Pagamentos e notificações fecham o ciclo: a multa pode ser resolvida pelo morador, "
             "com comprovante, e todos os eventos relevantes geram avisos no sistema.")

# ---------------------------------------------------------------------------
# 10 — Tecnologia e segurança
# ---------------------------------------------------------------------------
slide = new_slide()
header(slide, "TECNOLOGIA & SEGURANÇA", [
    [("Base sólida para", {"font": SERIF, "size": 31, "color": INK})],
    [("evoluir com segurança.", {"font": SERIF, "size": 31, "italic": True, "color": GOLD_DARK})],
], "Stack moderna, dados leves e controle de acesso desde a primeira linha.")

stack = [
    ("Next.js 15", "Interface rápida e renderizada no servidor."),
    ("TypeScript", "Menos erros e regras de negócio tipadas."),
    ("Tailwind CSS", "Visual consistente e responsivo."),
    ("SQLite", "Banco local simples, leve e auditável."),
]
for index, (title, text) in enumerate(stack):
    col = index % 2
    row = index // 2
    x = 0.65 + col * 2.72
    y = 2.45 + row * 1.55
    add_rect(slide, x, y, 2.52, 1.35, CARD, border=LINE, radius=0.05)
    add_text(slide, x + 0.22, y + 0.2, 2.1, 0.28,
             [[(title, {"size": 11.5, "bold": True, "color": INK})]])
    add_text(slide, x + 0.22, y + 0.55, 2.1, 0.65,
             [[(text, {"size": 8.2, "color": MUTED})]], line_spacing=1.15)

add_rect(slide, 6.55, 1.55, 6.13, 4.95, INK, radius=0.05)
add_text(slide, 6.85, 1.78, 5.5, 0.25,
         [[("SEGURANÇA", {"size": 8, "bold": True, "color": GOLD_LIGHT})]])
security = [
    ("Senhas com scrypt + salt", "Nada de senha em texto puro no banco."),
    ("Sessão em cookie httpOnly", "Token aleatório; só o hash fica armazenado."),
    ("Permissões por papel", "Verificadas em cada API, serviço e página."),
    ("Trilha de multas e pagamentos", "Valores, datas, método e comprovante registrados."),
    ("76 testes automatizados", "As regras críticas são testadas via API."),
]
for index, (title, text) in enumerate(security):
    y = 2.25 + index * 0.78
    add_circle(slide, 6.85, y + 0.02, 0.22, GOLD_PALE)
    add_text(slide, 6.85, y + 0.02, 0.22, 0.22,
             [[("✓", {"size": 8, "bold": True, "color": GOLD_DARK})]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 7.2, y - 0.02, 4.9, 0.25,
             [[(title, {"size": 10.3, "bold": True, "color": CREAM})]])
    add_text(slide, 7.2, y + 0.24, 4.9, 0.4,
             [[(text, {"size": 8, "color": RGBColor(0xB9, 0xCF, 0xC4)})]])

add_text(slide, 0.65, 5.65, 5.5, 0.7,
         [[("Responsivo no celular e no computador.", {"size": 9, "bold": True, "color": INK})],
          [("A mesma base atende morador, síndico e portaria.", {"size": 8, "color": MUTED})]],
         space_after=2)
add_line(slide, 6.85, 5.92, 5.5, RGBColor(0x3A, 0x55, 0x4B))
add_text(slide, 6.85, 6.02, 5.5, 0.35,
         [[("Dados do condomínio no SQLite; backups podem ser automatizados na implantação.",
            {"size": 7.5, "color": RGBColor(0xB9, 0xCF, 0xC4)})]], line_spacing=1.1)
footer(slide, 10)
notes(slide, "A stack é intencionalmente enxuta: Next.js e SQLite entregam velocidade e baixo custo "
             "de operação, com segurança de senhas, sessão e permissões desde o início.")

# ---------------------------------------------------------------------------
# 11 — Benefícios
# ---------------------------------------------------------------------------
slide = new_slide()
header(slide, "BENEFÍCIOS", [
    [("Menos conflito.", {"font": SERIF, "size": 31, "color": INK})],
    [("Mais convivência.", {"font": SERIF, "size": 31, "italic": True, "color": GOLD_DARK})],
], "O ganho não é só tecnológico: é de transparência, previsibilidade e tempo para todos.")

benefits = [
    ("Transparência", "Todos veem a mesma agenda e o mesmo histórico."),
    ("Regras claras", "Critérios objetivos para reservar, cancelar e pagar."),
    ("Prestação de contas", "Multas e pagamentos com registro e comprovante."),
    ("Pronto para crescer", "Novas áreas, regras e moradores sem refazer o sistema."),
]
for index, (title, text) in enumerate(benefits):
    x = 0.65 + index * 3.08
    add_rect(slide, x, 2.35, 2.9, 2.3, CARD, border=LINE, radius=0.05)
    add_text(slide, x + 0.25, 2.58, 0.8, 0.5,
             [[(f"{index + 1:02d}", {"font": SERIF, "size": 20, "color": GOLD})]])
    add_text(slide, x + 0.25, 3.12, 2.4, 0.3,
             [[(title, {"size": 12, "bold": True, "color": INK})]])
    add_text(slide, x + 0.25, 3.5, 2.4, 0.9,
             [[(text, {"size": 8.7, "color": MUTED})]], line_spacing=1.2)

add_rect(slide, 0.65, 5.15, 12.03, 1.2, INK, radius=0.05)
add_text(slide, 0.95, 5.3, 8.6, 0.42,
         [[("“A convivência começa quando a informação é a mesma para todos.”",
            {"font": SERIF, "size": 15, "italic": True, "color": CREAM})]])
add_text(slide, 0.95, 5.78, 8.6, 0.32,
         [[("Condomínio São Cristóvão · Área Comum", {"size": 8, "color": GOLD_LIGHT})]])
add_pill(slide, 10.0, 5.55, 2.35, 0.3, "4 áreas · 3 perfis · 1 sistema", GOLD_PALE, GOLD_DARK, size=7)
footer(slide, 11)
notes(slide, "Os benefícios vão além da reserva: o condomínio ganha transparência, critérios claros, "
             "registro financeiro e uma base que pode crescer com novas áreas e usuários.")

# ---------------------------------------------------------------------------
# 12 — Próximos passos e acesso
# ---------------------------------------------------------------------------
slide = new_slide(INK)
slide.shapes.add_picture(str(ASSETS / "watermark.png"), Inches(9.8), Inches(-1.4), Inches(4.6), Inches(4.6))
add_text(slide, 0.65, 0.5, 8.5, 0.25,
         [[("ENCAMINHAMENTO", {"size": 8.5, "bold": True, "color": GOLD_LIGHT})]])
add_text(slide, 0.65, 0.83, 10.5, 1.2,
         [[("Próximo passo:", {"font": SERIF, "size": 30, "color": CREAM}),
           (" aprovação e implantação.", {"font": SERIF, "size": 30, "italic": True, "color": GOLD_LIGHT})]])
add_text(slide, 0.65, 1.75, 10.5, 0.4,
         [[("Uma proposta pronta para ser apresentada e aprovada em assembleia.", {"size": 10, "color": RGBColor(0xB9, 0xCF, 0xC4)})]])

steps = [
    ("01", "Aprovação", "Apresentação e aprovação em assembleia."),
    ("02", "Configuração", "Áreas, regras, moradores e contas oficiais."),
    ("03", "Piloto", "Teste com um grupo e ajuste dos fluxos."),
    ("04", "Lançamento", "Treinamento da portaria e comunicação."),
]
for index, (number, title, text) in enumerate(steps):
    x = 0.65 + index * 3.06
    add_text(slide, x, 2.38, 0.7, 0.3,
             [[(number, {"font": SERIF, "size": 17, "color": GOLD_LIGHT})]])
    add_text(slide, x, 2.78, 2.7, 0.28,
             [[(title, {"size": 11, "bold": True, "color": CREAM})]])
    add_text(slide, x, 3.1, 2.65, 0.55,
             [[(text, {"size": 8, "color": RGBColor(0xB9, 0xCF, 0xC4)})]], line_spacing=1.15)
    if index < 3:
        add_line(slide, x + 0.62, 2.44, 2.28, RGBColor(0x3A, 0x55, 0x4B), height=0.008)

add_rect(slide, 0.65, 4.2, 12.03, 2.2, CARD, radius=0.05)
add_text(slide, 0.95, 4.36, 6.0, 0.25,
         [[("SISTEMA NO AR — ACESSE AGORA", {"size": 8, "bold": True, "color": GOLD_DARK})]])
add_text(slide, 0.95, 4.63, 3.8, 0.72,
         [[("area-comum.onrender.com", {"size": 14.5, "bold": True, "color": INK})],
          [("Link público, gratuito e com HTTPS.", {"size": 8.5, "color": MUTED})]],
         space_after=2)
add_text(slide, 0.95, 5.3, 3.85, 0.85,
         [[("Abra em qualquer computador ou celular, sem instalar nada nem usar a rede do condomínio. "
            "Na primeira visita após um período sem uso, o servidor gratuito pode levar até 1 minuto para responder.",
            {"size": 7.3, "color": MUTED})]], line_spacing=1.2)
add_text(slide, 0.95, 6.08, 3.85, 0.24,
         [[("Demonstração local: http://127.0.0.1:3000 (neste computador)", {"size": 6.6, "color": MUTED})]])

# QR Code do link público
if (ASSETS / "qr-acesso.png").exists():
    slide.shapes.add_picture(str(ASSETS / "qr-acesso.png"), Inches(5.02), Inches(4.6), Inches(1.4), Inches(1.4))
    add_text(slide, 5.02, 6.06, 1.4, 0.24,
             [[("Aponte a câmera", {"size": 6.5, "bold": True, "color": GOLD_DARK})]], align=PP_ALIGN.CENTER)

add_line(slide, 7.25, 4.45, 0.012, LINE, height=1.75)
add_text(slide, 7.55, 4.38, 5.0, 0.25,
         [[("CONTAS DE TESTE", {"size": 8, "bold": True, "color": GOLD_DARK})]])
accounts = [
    ("Síndico", "sindico@areacomum.app · Sindico@123"),
    ("Portaria", "portaria@areacomum.app · Portaria@123"),
    ("Morador", "ana@areacomum.app · Ana@1234"),
]
for index, (role, credential) in enumerate(accounts):
    y = 4.72 + index * 0.42
    add_text(slide, 7.55, y, 0.85, 0.25,
             [[(role, {"size": 8, "bold": True, "color": INK})]])
    add_text(slide, 8.42, y, 4.0, 0.25,
             [[(credential, {"size": 8, "color": MUTED})]])

add_text(slide, 0.65, 6.78, 8.0, 0.3,
         [[("Sistema publicado, com testes automatizados de regras e capturas do site em produção.", {"size": 7.5, "color": RGBColor(0x9F, 0xB3, 0xA8)})]])
add_text(slide, 10.5, 6.78, 2.2, 0.3,
         [[("Outubro de 2026", {"size": 7.5, "color": RGBColor(0x9F, 0xB3, 0xA8)})]], align=PP_ALIGN.RIGHT)
notes(slide, "Encerramento com o caminho de implantação e os acessos de demonstração. "
             "O sistema já está publicado: area-comum.onrender.com (link público e gratuito, com HTTPS). "
             "Mostre o QR Code para a assembleia abrir no celular na hora, ou use as contas de teste ao lado. "
             "A aprovação pode acontecer em assembleia e o piloto valida a operação antes do lançamento.")

prs.save(PPTX)
print("PPTX gerado:", PPTX)
print("Slides:", len(prs.slides))
